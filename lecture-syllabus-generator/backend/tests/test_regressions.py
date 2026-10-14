from io import BytesIO

import pytest
from docx import Document

from app.db import SessionLocal
from app.errors import ModelTimeoutError
from app.models import KnowledgePoint, Page, Task
from app.services.extractor import ExtractedKnowledgePoint, Extractor
from app.services.exporter import exporter
from app.services.outline_service import outline_service
from app.services.pipeline import pipeline
from app.services.recognizer import MockRecognizer, RecognitionResult


def upload(client, count=1):
    response = client.post('/api/tasks', files=[
        ('files', (f'page{i}.png', b'\x89PNG\r\n\x1a\n', 'image/png'))
        for i in range(count)
    ])
    assert response.status_code == 201
    return response.json()['task_id']


@pytest.mark.parametrize('tag,color', [('课堂原话', '333333'), ('AI推断', '0066CC'), ('来源待核对', None)])
def test_source_tags_survive_parse_and_word_export(tag, color):
    markdown = f'# 课程大纲\n## 第一章\n- 定义【集合】（来源：P1,P2）【{tag}】'
    point = outline_service.parse_markdown(markdown)[0]['points'][0]
    assert point == {'content': '定义【集合】', 'source_pages': 'P1,P2', 'source_type': tag}
    paragraph = Document(BytesIO(exporter.export(markdown))).paragraphs[-1]
    assert paragraph.text == f'定义【集合】（来源：P1,P2） 【{tag}】'
    actual_color = paragraph.runs[0].font.color.rgb
    assert (str(actual_color) if actual_color is not None else None) == color


def test_brackets_in_content_are_not_mistaken_for_source_tags():
    markdown = '# 课程大纲\n## 第一章\n- 术语【集合】'
    point = outline_service.parse_markdown(markdown)[0]['points'][0]
    assert point['content'] == '术语【集合】'
    assert point['source_type'] == ''
    assert Document(BytesIO(exporter.export(markdown))).paragraphs[-1].text == '术语【集合】'


def test_recognize_then_generate_preserves_source_types_after_reload(client, monkeypatch):
    task_id = upload(client)
    assert client.post(f'/api/tasks/{task_id}/recognize').status_code == 200

    def unexpected_recognition(*args):
        raise AssertionError('Saved recognition must be reused')

    monkeypatch.setattr(pipeline.recognizer, 'recognize', unexpected_recognition)
    for endpoint in ['pipeline', 'regenerate']:
        assert client.post(f'/api/tasks/{task_id}/{endpoint}').status_code == 200
        markdown = client.get(f'/api/tasks/{task_id}/outline').json()['markdown']
        assert '数据模型的三要素：数据结构、数据操作、数据约束（来源：P1）【AI推断】' in markdown
        assert '关系数据库基本概念（来源：P1）【课堂原话】' in markdown
    with SessionLocal() as db:
        point = db.query(KnowledgePoint).filter_by(task_id=task_id, content='数据模型的三要素：数据结构、数据操作、数据约束').one()
        assert point.source_type == 'AI推断'


def test_legacy_text_does_not_invent_original_source(client):
    task_id = upload(client)
    with SessionLocal() as db:
        page = db.query(Page).filter_by(task_id=task_id).one()
        page.recognized_text = '历史标题\n- 没有保存来源类型的内容'
        db.commit()
    assert client.post(f'/api/tasks/{task_id}/pipeline').status_code == 200
    markdown = client.get(f'/api/tasks/{task_id}/outline').json()['markdown']
    assert '【来源待核对】' in markdown
    assert '【课堂原话】' not in markdown


def test_text_only_model_response_keeps_title_when_saved_structure_is_empty(client, monkeypatch):
    task_id = upload(client)
    monkeypatch.setattr(pipeline.recognizer, 'recognize', lambda *args: RecognitionResult(
        text='真实课程标题\n- 图片识别出的知识点',
        structured={'title': '页面1', 'points': []},
    ))
    assert client.post(f'/api/tasks/{task_id}/recognize').status_code == 200
    assert client.post(f'/api/tasks/{task_id}/pipeline').status_code == 200
    markdown = client.get(f'/api/tasks/{task_id}/outline').json()['markdown']
    assert '## 一、真实课程标题' in markdown
    assert '图片识别出的知识点（来源：P1）【来源待核对】' in markdown


def test_dedup_keeps_different_source_types_separate():
    points = [
        ExtractedKnowledgePoint('同一知识点', 'concept', '课堂原话', 'P1', 1, '同一知识点'),
        ExtractedKnowledgePoint('同一知识点', 'concept', 'AI推断', 'P2', 2, '同一知识点'),
        ExtractedKnowledgePoint('同一知识点', 'concept', '课堂原话', 'P3', 3, '同一知识点'),
    ]
    result = Extractor().dedup_cross_page(points)
    assert {(p.source_type, p.source_pages) for p in result} == {('课堂原话', 'P1,P3'), ('AI推断', 'P2')}


def test_edit_reopen_and_download_word_have_same_content(client):
    task_id = upload(client)
    assert client.post(f'/api/tasks/{task_id}/pipeline').status_code == 200
    old_version = client.get(f'/api/tasks/{task_id}/outline').json()['version']
    edited = '# 我修改的课程大纲\n\n## 第一模块\n- 自己改过的定义（来源：P1）【课堂原话】\n- 需要核对的推断（来源：P1）【AI推断】\n\n## 第二模块\n补充复习说明'
    saved = client.put(f'/api/tasks/{task_id}/outline', json={'markdown': edited})
    assert saved.status_code == 200
    assert saved.json()['version'] == old_version + 1
    reopened = client.get(f'/api/tasks/{task_id}/outline').json()
    assert reopened['markdown'] == edited
    response = client.get(f'/api/tasks/{task_id}/export')
    assert response.status_code == 200
    assert '.docx' in response.headers['content-disposition']
    paragraphs = Document(BytesIO(response.content)).paragraphs
    assert [p.text for p in paragraphs] == [
        '我修改的课程大纲', '第一模块',
        '自己改过的定义（来源：P1） 【课堂原话】',
        '需要核对的推断（来源：P1） 【AI推断】', '第二模块', '补充复习说明',
    ]
    assert str(paragraphs[3].runs[0].font.color.rgb) == '0066CC'


@pytest.mark.parametrize('method,path', [
    ('get', '/api/tasks/999999'), ('get', '/api/tasks/999999/outline'),
    ('get', '/api/tasks/999999/export'), ('post', '/api/tasks/999999/pipeline'),
])
def test_missing_task_returns_real_http_error_contract(client, method, path):
    response = getattr(client, method)(path)
    assert response.status_code == 404
    assert response.json() == {'code': 'TASK_NOT_FOUND', 'message': '任务不存在。', 'retryable': False}


def test_invalid_upload_has_clear_error_and_creates_no_task(client):
    response = client.post('/api/tasks', files={'files': ('bad.txt', b'bad', 'text/plain')})
    assert response.status_code == 400
    assert response.json()['code'] == 'FILE_FORMAT_ERROR'
    assert response.json()['retryable'] is False
    assert client.get('/api/tasks').json()['tasks'] == []


def test_export_before_generation_returns_error_not_a_broken_docx(client):
    task_id = upload(client)
    response = client.get(f'/api/tasks/{task_id}/export')
    assert response.status_code == 500
    assert response.json()['code'] == 'PIPELINE_ERROR'
    assert response.json()['retryable'] is True


def test_direct_pipeline_continues_after_one_page_timeout(client, monkeypatch):
    task_id = upload(client, count=2)
    mock = MockRecognizer()

    def recognize(page_no, path):
        if page_no == 1:
            raise ModelTimeoutError()
        return mock.recognize(page_no, path)

    monkeypatch.setattr(pipeline.recognizer, 'recognize', recognize)
    response = client.post(f'/api/tasks/{task_id}/pipeline')
    assert response.status_code == 200
    assert response.json()['status'] == 'partial'
    pages = client.get(f'/api/tasks/{task_id}/pages').json()['pages']
    assert pages[0]['failed_flag'] is True
    assert '超时' in pages[0]['error_msg']
    assert pages[1]['failed_flag'] is False
    assert '范式与函数依赖' in client.get(f'/api/tasks/{task_id}/outline').json()['markdown']


def test_all_page_failures_do_not_create_an_empty_success_outline(client, monkeypatch):
    task_id = upload(client)

    def timeout(*args):
        raise ModelTimeoutError()

    monkeypatch.setattr(pipeline.recognizer, 'recognize', timeout)
    assert client.post(f'/api/tasks/{task_id}/recognize').json()['status'] == 'failed'
    response = client.post(f'/api/tasks/{task_id}/pipeline')
    assert response.status_code == 500
    assert response.json()['code'] == 'PIPELINE_ERROR'
    assert response.json()['retryable'] is True
    assert client.get(f'/api/tasks/{task_id}').json()['status'] == 'failed'
    assert client.get(f'/api/tasks/{task_id}/outline').status_code == 500


def test_retry_status_accounts_for_other_failed_pages(client, monkeypatch):
    task_id = upload(client, count=2)

    def timeout(*args):
        raise ModelTimeoutError()

    monkeypatch.setattr(pipeline.recognizer, 'recognize', timeout)
    client.post(f'/api/tasks/{task_id}/recognize')
    monkeypatch.setattr(pipeline.recognizer, 'recognize', MockRecognizer().recognize)
    response = client.post(f'/api/tasks/{task_id}/recognize?page_no=1')
    assert response.json()['status'] == 'partial'
    pages = client.get(f'/api/tasks/{task_id}/pages').json()['pages']
    assert pages[0]['failed_flag'] is False
    assert pages[0]['error_msg'] is None
    assert pages[1]['failed_flag'] is True
    assert client.post(f'/api/tasks/{task_id}/recognize?page_no=2').json()['status'] == 'success'


def test_unexpected_generation_error_is_safe_and_not_stuck_running(client, monkeypatch):
    task_id = upload(client)

    def crash(*args):
        raise RuntimeError('private-debug-value')

    monkeypatch.setattr(outline_service, 'render_markdown', crash)
    response = client.post(f'/api/tasks/{task_id}/pipeline')
    assert response.status_code == 500
    assert response.json() == {'code': 'INTERNAL_ERROR', 'message': '服务器内部错误。', 'retryable': False}
    status = client.get(f'/api/tasks/{task_id}').json()
    assert status['status'] == 'failed'
    assert 'private-debug-value' not in str(status)


def test_failed_regeneration_keeps_saved_outline_and_knowledge_points(client, monkeypatch):
    task_id = upload(client)
    client.post(f'/api/tasks/{task_id}/pipeline')
    edited = '# 我的笔记\n## 已校对\n- 保留这条修改（来源：P1）【课堂原话】'
    client.put(f'/api/tasks/{task_id}/outline', json={'markdown': edited})
    with SessionLocal() as db:
        original_points = [(p.id, p.content, p.source_type) for p in db.query(KnowledgePoint).filter_by(task_id=task_id).order_by(KnowledgePoint.id)]
        page = db.query(Page).filter_by(task_id=task_id).one()
        page.recognized_structured = {'title': '新标题', 'points': [
            {'content': '新识别结果', 'kp_type': 'concept', 'source_type': 'AI推断'},
        ]}
        db.commit()

    def crash(*args):
        raise RuntimeError('generation interrupted')

    monkeypatch.setattr(pipeline.clusterer, 'cluster', crash)
    assert client.post(f'/api/tasks/{task_id}/regenerate').status_code == 500
    assert client.get(f'/api/tasks/{task_id}/outline').json()['markdown'] == edited
    with SessionLocal() as db:
        assert [(p.id, p.content, p.source_type) for p in db.query(KnowledgePoint).filter_by(task_id=task_id).order_by(KnowledgePoint.id)] == original_points


def test_failed_retry_does_not_reuse_stale_recognition(client, monkeypatch):
    task_id = upload(client)
    client.post(f'/api/tasks/{task_id}/recognize')

    def timeout(*args):
        raise ModelTimeoutError()

    monkeypatch.setattr(pipeline.recognizer, 'recognize', timeout)
    assert client.post(f'/api/tasks/{task_id}/recognize?page_no=1').json()['status'] == 'failed'
    with SessionLocal() as db:
        page = db.query(Page).filter_by(task_id=task_id).one()
        assert page.recognized_text is None
        assert page.recognized_structured is None
    assert client.post(f'/api/tasks/{task_id}/pipeline').status_code == 500


def test_retrying_missing_page_does_not_change_task_status(client):
    task_id = upload(client)
    before = client.get(f'/api/tasks/{task_id}').json()['status']
    response = client.post(f'/api/tasks/{task_id}/recognize?page_no=999')
    assert response.status_code == 404
    assert response.json()['code'] == 'PAGE_NOT_FOUND'
    assert client.get(f'/api/tasks/{task_id}').json()['status'] == before


def test_single_page_recognition_does_not_mark_unprocessed_batch_successful(client):
    task_id = upload(client, count=2)
    assert client.post(f'/api/tasks/{task_id}/recognize?page_no=1').json()['status'] == 'pending'


def test_empty_task_pipeline_has_clear_error_and_failed_status(client):
    with SessionLocal() as db:
        task = Task(title='空任务', page_count=0)
        db.add(task)
        db.commit()
        task_id = task.id
    response = client.post(f'/api/tasks/{task_id}/pipeline')
    assert response.status_code == 500
    assert response.json()['code'] == 'PIPELINE_ERROR'
    assert client.get(f'/api/tasks/{task_id}').json()['status'] == 'failed'
