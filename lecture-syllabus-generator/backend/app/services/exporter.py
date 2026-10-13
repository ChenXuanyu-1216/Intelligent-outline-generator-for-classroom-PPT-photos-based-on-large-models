from abc import ABC, abstractmethod
from io import BytesIO
from docx import Document
from docx.shared import Pt, RGBColor
from re import sub


class Exporter(ABC):
    @abstractmethod
    def export(self, markdown: str) -> bytes:
        pass


class DocxExporter(Exporter):
    def export(self, markdown: str) -> bytes:
        doc = Document()

        style = doc.styles["Normal"]
        style.font.name = "宋体"
        style.font.size = Pt(12)

        lines = markdown.strip().split("\n")
        for line in lines:
            stripped = line.strip()
            if not stripped:
                continue
            if stripped.startswith("# "):
                heading = stripped[2:].strip()
                h = doc.add_heading(heading, level=0)
            elif stripped.startswith("## "):
                heading = stripped[3:].strip()
                h = doc.add_heading(heading, level=1)
            elif stripped.startswith("### "):
                heading = stripped[4:].strip()
                h = doc.add_heading(heading, level=2)
            elif stripped.startswith("- "):
                content = stripped[2:].strip()
                self._add_bullet_with_source(doc, content)
            else:
                doc.add_paragraph(stripped)

        buf = BytesIO()
        doc.save(buf)
        return buf.getvalue()

    def _add_bullet_with_source(self, doc, content: str):
        source_type = ""
        if "【" in content and "】" in content:
            s = content.rfind("【")
            e = content.rfind("】")
            if e > s:
                source_type = content[s + 2:e]
                content = content[:s] + content[e + 1:]

        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(content.strip())
        if source_type == "AI推断":
            run.font.color.rgb = RGBColor(0x00, 0x66, 0xCC)
        elif source_type == "课堂原话":
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

        if source_type:
            tag_run = p.add_run(f" 【{source_type}】")
            tag_run.font.size = Pt(10)
            if source_type == "AI推断":
                tag_run.font.color.rgb = RGBColor(0x00, 0x66, 0xCC)
            else:
                tag_run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)


exporter = DocxExporter()